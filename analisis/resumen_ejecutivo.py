# -*- coding: utf-8 -*-
"""Resumen ejecutivo de cobertura IPERC - una serie, barras de datos, estado con etiqueta."""
import json, datetime, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import DataBarRule
from openpyxl.utils import get_column_letter

# paleta: serie unica azul; estado con etiqueta textual, nunca color solo
SERIE="2A78D6"; GOOD="0CA30C"; WARN="FAB219"; CRIT="D03B3B"
TINTA="0B0B0B"; TINTA2="52514E"; AZUL="1F4E79"; SUP="FCFCFB"
_t=Side(style="thin",color="D9D9D9"); BOR=Border(left=_t,right=_t,top=_t,bottom=_t)
CEN=Alignment(horizontal="center",vertical="center",wrap_text=True)
IZQ=Alignment(horizontal="left",vertical="center",wrap_text=True)
F=lambda sz=10,b=False,col=TINTA: Font(name="Calibri",size=sz,bold=b,color=col)
fill=lambda c: PatternFill("solid",fgColor=c)

src=openpyxl.load_workbook("Cobertura_IPERC_Administrativos_CAR.xlsx")
areas=[]
for r in src["2. Resumen por área"].iter_rows(min_row=5,values_only=True):
    if not r[0]: continue
    if str(r[0]).startswith("TOTAL"): tot=r; continue
    areas.append(r)
areas.sort(key=lambda x:(-x[4],x[0]))
brechas=[r for r in src["4. Brechas"].iter_rows(min_row=5,values_only=True) if r and r[0]]
matrices=sum(1 for r in src["6. Trazabilidad"].iter_rows(min_row=5,values_only=True)
             if r and r[0] and r[0]!="TOTAL")
eqd=json.load(open("analisis/equivalencias.json",encoding="utf-8"))

wb=openpyxl.Workbook(); ws=wb.active; ws.title="Resumen ejecutivo"
ws.sheet_view.showGridLines=False
for col,w in zip("ABCDEFGHI",[30,10,11,11,13,10,11,13,16]):
    ws.column_dimensions[col].width=w

ws.merge_cells("A1:I1"); c=ws["A1"]
c.value="COBERTURA DE LA IDENTIFICACIÓN DE PELIGROS Y EVALUACIÓN DE RIESGOS"
c.font=F(14,True,"FFFFFF"); c.fill=fill(AZUL); c.alignment=CEN
ws.row_dimensions[1].height=30
ws.merge_cells("A2:I2"); c=ws["A2"]
c.value=("Personal administrativo  ·  Plantaciones de Balsa Plantabal S.A. / 3A Core Materials  ·  "
         f"Head count de octubre de 2026  ·  Corte al {datetime.date.today().strftime('%d/%m/%Y')}")
c.font=Font(name="Calibri",size=9,italic=True,color=TINTA2); c.alignment=CEN
ws.row_dimensions[2].height=18

# indicadores principales
tiles=[("A4","C6",f"{tot[4]}%","de los cargos con matriz",f"{tot[2]} de {tot[1]} cargos"),
       ("D4","F6",f"{round(100*tot[6]/tot[5])}%","del personal cubierto",f"{tot[6]} de {tot[5]} personas"),
       ("G4","H6",str(matrices),"matrices IPERC","analizadas"),
       ("I4","I6",str(len(brechas)),"brechas","por cerrar")]
for ini,fin,num,et1,et2 in tiles:
    ws.merge_cells(f"{ini}:{fin}")
    cc=ws[ini]; cc.value=num; cc.font=F(26,True,AZUL); cc.alignment=Alignment(horizontal="center",vertical="center")
    cc.fill=fill("EFF4FA"); cc.border=BOR
    col=ini[0]; row=7
    ws.merge_cells(f"{col}{row}:{fin[0]}{row}")
    d=ws[f"{col}{row}"]; d.value=f"{et1}\n{et2}"; d.font=F(9,False,TINTA2)
    d.alignment=Alignment(horizontal="center",vertical="top",wrap_text=True)
ws.row_dimensions[7].height=26

# tabla por area
hdr=["Área","Cargos","Analizados","Sin analizar","% de cargos","Personas","Cubiertas","% de personas","Estado"]
for j,t in enumerate(hdr,1):
    cc=ws.cell(9,j,t); cc.font=F(10,True,"FFFFFF"); cc.fill=fill(AZUL); cc.alignment=CEN; cc.border=BOR
ws.row_dimensions[9].height=30
fila=10
for a in areas:
    pp=round(100*a[6]/a[5]) if a[5] else 0
    est="Conforme" if a[3]==0 else ("Crítico" if a[4]<60 else "Con brechas")
    for j,v in enumerate([a[0],a[1],a[2],a[3],a[4],a[5],a[6],pp,est],1):
        cc=ws.cell(fila,j,v); cc.border=BOR; cc.font=F(10)
        cc.alignment=IZQ if j in (1,9) else CEN
        if j in (5,8): cc.number_format='0"%"'
        if fila%2==0: cc.fill=fill("FAFAF9")
    ws.cell(fila,9).font=F(10,True,{"Conforme":GOOD,"Con brechas":WARN,"Crítico":CRIT}[est])
    fila+=1
ult=fila-1
ppt=round(100*tot[6]/tot[5])
for j,v in enumerate(["TOTAL GENERAL",tot[1],tot[2],tot[3],tot[4],tot[5],tot[6],ppt,""],1):
    cc=ws.cell(fila,j,v); cc.font=F(10,True); cc.fill=fill("D9E2F3"); cc.border=BOR
    cc.alignment=IZQ if j in (1,9) else CEN
    if j in (5,8): cc.number_format='0"%"' 
# barras de datos: una sola serie, sin ejes, el valor queda visible
for col in ("E","H"):
    ws.conditional_formatting.add(f"{col}10:{col}{ult}",
        DataBarRule(start_type="num",start_value=0,end_type="num",end_value=100,
                    color=SERIE,showValue=True))
ws.freeze_panes="A10"

nota=fila+2
ws.merge_cells(f"A{nota}:I{nota}")
c=ws[f"A{nota}"]
c.value=("Alcance: personal administrativo de Quevedo y Samborondón. Se excluyen el área de Silvicultura, "
         "los cargos operativos y el personal con localidad principal High Point. El estado 'Crítico' marca "
         "las áreas con menos del 60% de sus cargos evaluados.")
c.font=F(9,False,TINTA2); c.alignment=IZQ
ws.row_dimensions[nota].height=28

# ── brechas priorizadas
ws2=wb.create_sheet("Brechas priorizadas"); ws2.sheet_view.showGridLines=False
for col,w in zip("ABCDEF",[30,44,12,14,16,34]): ws2.column_dimensions[col].width=w
ws2.merge_cells("A1:F1"); c=ws2["A1"]; c.value="BRECHAS POR CERRAR, ORDENADAS POR PERSONAS EXPUESTAS"
c.font=F(13,True,"FFFFFF"); c.fill=fill(AZUL); c.alignment=CEN; ws2.row_dimensions[1].height=26
hdr2=["Área","Cargo sin evaluar","N° personas","% del total sin cubrir","% acumulado","Acción requerida"]
for j,t in enumerate(hdr2,1):
    cc=ws2.cell(3,j,t); cc.font=F(10,True,"FFFFFF"); cc.fill=fill(AZUL); cc.alignment=CEN; cc.border=BOR
ws2.row_dimensions[3].height=30
br=sorted(brechas,key=lambda x:-x[2]); sincub=sum(b[2] for b in br); acum=0
for i,b in enumerate(br,4):
    acum+=b[2]
    for j,v in enumerate([b[0],b[1],b[2],round(100*b[2]/sincub),round(100*acum/sincub),
                          "Elaborar matriz IPERC para el puesto"],1):
        cc=ws2.cell(i,j,v); cc.border=BOR; cc.font=F(10); cc.alignment=IZQ if j in (1,2,6) else CEN
        if j in (4,5): cc.number_format='0"%"'
        if i%2==0: cc.fill=fill("FAFAF9")
    ws2.cell(i,2).font=F(10,True,CRIT)
ws2.conditional_formatting.add(f"D4:D{3+len(br)}",
    DataBarRule(start_type="num",start_value=0,end_type="num",end_value=100,color=SERIE,showValue=True))
ws2.freeze_panes="A4"

# ── acciones correctivas
ws3=wb.create_sheet("Acciones correctivas"); ws3.sheet_view.showGridLines=False
for col,w in zip("ABCD",[40,34,18,56]): ws3.column_dimensions[col].width=w
ws3.merge_cells("A1:D1"); c=ws3["A1"]; c.value="ESTADO DE LAS ACCIONES CORRECTIVAS DOCUMENTALES"
c.font=F(13,True,"FFFFFF"); c.fill=fill(AZUL); c.alignment=CEN; ws3.row_dimensions[1].height=26
for j,t in enumerate(["Hallazgo","Tipo","Estado","Detalle"],1):
    cc=ws3.cell(3,j,t); cc.font=F(10,True,"FFFFFF"); cc.fill=fill(AZUL); cc.alignment=CEN; cc.border=BOR
acc=[]
for k,v in sorted(eqd.get("cargos_regularizados",{}).items()):
    acc.append([k,"Cargo que solo existía en la matriz","CERRADA",v.get("estado","")])
for k,v in sorted(eqd.get("cargos_obsoletos_a_eliminar",{}).items()):
    acc.append([k,"Cargo obsoleto en la matriz","ABIERTA",v.get("accion","")])
acc.append([f"{len(eqd['equivalencias'])} equivalencias de denominación","Nomenclatura divergente","ABIERTA",
            "Unificar las denominaciones de las matrices según el head count, como fuente única"])
for i,a in enumerate(acc,4):
    for j,v in enumerate(a,1):
        cc=ws3.cell(i,j,v); cc.border=BOR; cc.font=F(10); cc.alignment=IZQ if j!=3 else CEN
        if i%2==0: cc.fill=fill("FAFAF9")
    ws3.cell(i,3).font=F(10,True,GOOD if a[2]=="CERRADA" else WARN)
ws3.freeze_panes="A4"

wb.save("Resumen_Ejecutivo_Cobertura_IPERC.xlsx")
print("Generado: Resumen_Ejecutivo_Cobertura_IPERC.xlsx")
print(f"  Áreas: {len(areas)} | Brechas: {len(br)} ({sincub} personas) | Acciones: {len(acc)}")
