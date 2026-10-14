# -*- coding: utf-8 -*-
"""
Cruce de cobertura IPERC a partir de la TABLA DINAMICA del head count.

Para los meses en que el head count llega solo como tabla dinamica (imagen o
copia de centro de costo + cargo) y no como archivo. El resultado se expresa en
cargos, no en personas: la tabla no trae dotacion ni area.

Entrada: un archivo de texto con una linea por par, separado por |
    1200 PLANTATION MANAGEMENT|COORDINADOR DE SEGURIDAD FISICA
    5800 GENERAL MANAGEMENT|GERENTE GENERAL

Uso:
    python3 cruce_tabla.py --tabla hc_tabla_octubre_2026.txt [--out archivo.xlsx]
"""
import argparse, json, os, sys, importlib.util
from collections import defaultdict, OrderedDict
import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))
_spec = importlib.util.spec_from_file_location("cm", os.path.join(BASE, "control_mensual.py"))
cm = importlib.util.module_from_spec(_spec); _spec.loader.exec_module(cm)
norm = cm.norm

# centros de costo fuera del alcance acordado, por su prefijo numerico
CC_FUERA = {
    "5007": "Fuera de alcance - High Point (Product Manager)",
    "5033": "Fuera de alcance - High Point (PGS-CMBA Supply Chain)",
    "5927": "Fuera de alcance - High Point (R&D)",
    "5937": "Fuera de alcance - High Point (Certification MX)",
    "4207": "Fuera de alcance - Santo Domingo",
    "4507": "Fuera de alcance - Santo Domingo",
}

def J(n): return json.load(open(os.path.join(BASE, n), encoding="utf-8"))

def main():
    ap = argparse.ArgumentParser(description="Cruce IPERC desde la tabla dinamica del head count")
    ap.add_argument("--tabla", required=True, help="Archivo de texto con lineas 'CENTRO DE COSTO|CARGO'")
    ap.add_argument("--out", default="Cruce_HC_vs_Matrices.xlsx")
    a = ap.parse_args()

    cat = J("catalogo_matrices.json")["matrices"]
    EQ = {}
    for k, v in J("equivalencias.json")["equivalencias"].items():
        d = v["cargo_hc"]; EQ[k] = [d] if isinstance(d, str) else list(d)
    ev = defaultdict(list)
    for m in cat:
        for pu, n in m["puestos"].items():
            k = norm(pu)
            for destino in EQ.get(k, [k]):
                ev[norm(destino)].append((m["codigo"], pu, n))
    OPER = {norm(c) for c in J("cargos_administrativos.json")["operativos"]}

    ruta = a.tabla if os.path.isabs(a.tabla) else os.path.join(BASE, a.tabla)
    if not os.path.exists(ruta): sys.exit(f"ERROR: no se encontro {ruta}")
    pares = [l.strip().split("|") for l in open(ruta, encoding="utf-8") if l.strip() and "|" in l]
    print(f"  Pares centro de costo-cargo leidos: {len(pares)}")

    dentro, fuera, filas = OrderedDict(), [], []
    for cc, ca in pares:
        cod = cc.split()[0]; cn = norm(ca)
        if cod in CC_FUERA:
            fuera.append([cc, ca, CC_FUERA[cod], "No se evalua en este control"]); continue
        if cn in OPER:
            fuera.append([cc, ca, "Cargo operativo", "Fuera del universo administrativo"]); continue
        dentro.setdefault(cn, {"nom": ca, "cc": []})["cc"].append(cc)
        e = ev.get(cn, [])
        filas.append([cc, ca, "ANALIZADO" if e else "SIN ANALIZAR",
                      " / ".join(sorted({x[0] for x in e})), " / ".join(sorted({x[1] for x in e})),
                      sum(x[2] for x in e),
                      "Denominación distinta en la matriz (equivalencia validada)"
                      if e and norm(e[0][1]) != cn else ""])

    wb = openpyxl.Workbook(); wb.remove(wb.active)
    cm.hoja(wb, "1. Alcance y origen", "ALCANCE, ORIGEN DE LOS DATOS Y ADVERTENCIAS",
            ["N°", "Punto", "Detalle"],
            [["1", "Origen del head count", f"Transcripción de la tabla dinámica del head count ({os.path.basename(ruta)}). NO proviene del archivo Excel."],
             ["2", "Limitación: dotación", "La tabla muestra centros de costo y cargos, pero no el número de personas por cargo. Este cruce se expresa en cargos, no en personas."],
             ["3", "Limitación: área", "La tabla no trae la columna Área. La pertenencia al alcance se determinó por centro de costo."],
             ["4", "Exclusión por localidad", "Quedan fuera los centros de costo de High Point (5007, 5033, 5927, 5937) y de Santo Domingo (4207, 4507)."],
             ["5", "Exclusión por tipo de cargo", "Quedan fuera los cargos clasificados como operativos en el catálogo de cargos."],
             ["6", "Verificación recomendada", "Para un consolidado completo, con dotación por cargo y comparación automática contra el corte anterior, procesar el archivo Excel del head count con control_mensual.py."]],
            [6, 32, 108])
    ws = cm.hoja(wb, "2. Cruce por centro de costo", "CARGOS DEL HEAD COUNT vs. MATRICES IPERC",
                 ["Centro de costo", "Cargo", "Estado", "Código de matriz",
                  "Denominación en la matriz", "Líneas de riesgo", "Observación"],
                 filas, [36, 44, 15, 16, 42, 11, 52])
    for i in range(5, 5 + len(filas)):
        c = ws.cell(i, 3)
        if c.value == "ANALIZADO": c.font = cm.FB(10, cm.VD)
        else: c.font = cm.FB(10, cm.RJ); c.fill = cm.fill("FCE4E4")
    br = [[" / ".join(v["cc"]), v["nom"], "Elaborar matriz IPERC para el puesto", "", ""]
          for k, v in sorted(dentro.items()) if k not in ev]
    ws = cm.hoja(wb, "3. Brechas", "CARGOS SIN EVALUACIÓN DE RIESGOS",
                 ["Centro de costo", "Cargo sin analizar", "Acción requerida", "Responsable", "Fecha compromiso"],
                 br, [44, 44, 40, 20, 16])
    for i in range(5, 5 + len(br)): ws.cell(i, 2).font = cm.FB(10, cm.RJ)
    cm.hoja(wb, "4. Excluidos por alcance", "REGISTROS DE LA TABLA QUE QUEDAN FUERA DEL UNIVERSO",
            ["Centro de costo", "Cargo", "Motivo de exclusión", "Tratamiento"],
            sorted(set(map(tuple, fuera))), [36, 46, 46, 36])
    tz = [[m["codigo"], m["proceso"], len(m["puestos"]), sum(m["puestos"].values()),
           " · ".join(f"{k} ({v})" for k, v in m["puestos"].items())]
          for m in sorted(cat, key=lambda x: x["codigo"])]
    cm.hoja(wb, "5. Trazabilidad", "MATRICES IPERC ANALIZADAS",
            ["Código", "Puesto / proceso evaluado", "N° de puestos", "Líneas de riesgo", "Puestos evaluados"],
            tz, [16, 34, 12, 13, 95])
    wb.save(a.out)
    cub = len(dentro) - len(br)
    print(f"  Archivo generado: {a.out}")
    print(f"  Cobertura: {cub}/{len(dentro)} cargos ({round(100*cub/len(dentro))}%)"
          f"  |  brechas: {len(br)}  |  excluidos: {len(set(map(tuple, fuera)))}\n")

if __name__ == "__main__":
    main()
