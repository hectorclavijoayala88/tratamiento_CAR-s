import openpyxl, os, json, sys, re
from collections import OrderedDict

def norm(s):
    s=" ".join(str(s or "").strip().upper().split())
    for a,b in zip("ÁÉÍÓÚÜÑ","AEIOUUN"): s=s.replace(a,b)
    return s

def extraer(path):
    wb=openpyxl.load_workbook(path, data_only=True)
    res=[]
    for ws in wb.worksheets:
        # metadatos de cabecera
        meta={}
        for row in ws.iter_rows(min_row=1,max_row=8,max_col=30):
            for c in row:
                v=norm(c.value)
                if v in ("PUESTO DE TRABAJO A EVALUAR","CODIGO","AREA","PROCESO","EMPRESA"):
                    # buscar el primer valor no vacío a la derecha
                    for cc in ws[c.row]:
                        if cc.column>c.column and cc.value not in (None,""):
                            meta[v]=str(cc.value).strip(); break
        codigo=None
        for row in ws.iter_rows(min_row=1,max_row=3,max_col=30):
            for c in row:
                if isinstance(c.value,str) and re.match(r'^GE-SO-[A-Z]+-D\d',c.value.strip()):
                    codigo=c.value.strip()
        # localizar fila de encabezado y columna PUESTO DE TRABAJO
        hrow=hcol=None
        for row in ws.iter_rows(min_row=1,max_row=20,max_col=30):
            for c in row:
                if norm(c.value)=="PUESTO DE TRABAJO":
                    hrow,hcol=c.row,c.column; break
            if hrow: break
        if not hrow:
            res.append({"hoja":ws.title,"error":"no se encontró columna PUESTO DE TRABAJO"}); continue
        # rellenar celdas combinadas verticalmente
        fill={}
        for mr in ws.merged_cells.ranges:
            if mr.min_col==hcol:
                v=ws.cell(mr.min_row,mr.min_col).value
                for r in range(mr.min_row,mr.max_row+1): fill[r]=v
        puestos=OrderedDict()
        actual=None; filas_datos=0
        for r in range(hrow+1, ws.max_row+1):
            v=fill.get(r, ws.cell(r,hcol).value)
            # fila con contenido real?
            fila_tiene=any(ws.cell(r,c).value not in (None,"") for c in range(1,min(ws.max_column,30)+1))
            if not fila_tiene: continue
            filas_datos+=1
            if v not in (None,"") and str(v).strip():
                actual=str(v).strip()
            if actual:
                puestos.setdefault(actual,0)
                puestos[actual]+=1
        res.append({"hoja":ws.title,"codigo":codigo,"meta":meta,
                    "fila_encabezado":hrow,"filas_datos":filas_datos,
                    "puestos":puestos})
    return res

if __name__=="__main__":
    salida={}
    for p in sys.argv[1:]:
        salida[os.path.basename(p)]=extraer(p)
    print(json.dumps(salida,ensure_ascii=False,indent=1))
